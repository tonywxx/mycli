#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Extract speaker notes (备注) from .pptx file(s) and write them to .txt files.

Modes:
  1) Single file:
       python3 extract_pptx_notes.py <input.pptx> [output.txt]
     If output.txt is omitted, it defaults to "<input_basename>_note.txt"
     in the same directory as the input file.

  2) Directory (auto) mode:
       python3 extract_pptx_notes.py -auto <dir> [--depth N] [-filename=<kw>]
       python3 extract_pptx_notes.py <dir>            # directory auto-detected
     Recursively finds every .pptx under <dir> and writes "<name>_note.txt"
     next to each one. Default descent is 3 levels; --depth 0 means unlimited.

  Filtering:
       -filename=<keyword>   Only process .pptx whose filename (basename, without
                             the .pptx extension) contains <keyword>. Case-sensitive
                             substring match. Without this option, all .pptx are
                             processed. Accepts "-filename <keyword>" too.

Requires: python-pptx  (pip install python-pptx)
"""

import sys
import os
from pptx import Presentation


def extract_notes(pptx_path):
    """Return a list of (slide_no, notes_text) tuples for slides that have notes."""
    prs = Presentation(pptx_path)
    results = []
    for i, slide in enumerate(prs.slides, start=1):
        notes_text = ""
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text
        if notes_text and notes_text.strip():
            results.append((i, notes_text.strip()))
    return results


def write_notes(notes, out_path):
    """Write the extracted notes to out_path (UTF-8)."""
    with open(out_path, "w", encoding="utf-8") as f:
        if not notes:
            f.write("(No speaker notes found in this presentation.)\n")
        for slide_no, text in notes:
            f.write("===== 第 %d 页 备注 =====\n" % slide_no)
            f.write(text)
            f.write("\n\n")


def process_file(pptx_path, out_path=None):
    """Extract notes from a single .pptx. Returns the output path."""
    if out_path is None:
        base = os.path.splitext(pptx_path)[0]
        out_path = base + "_note.txt"
    notes = extract_notes(pptx_path)
    write_notes(notes, out_path)
    print("Extracted %d slide(s) with notes -> %s" % (len(notes), out_path))
    return out_path


def filename_matches(pptx_path, keyword):
    """Return True if keyword is None or the pptx basename contains keyword.

    The match is a case-sensitive substring test on the filename WITHOUT the
    .pptx extension (so "-filename=note" matches "foo note.pptx" but a keyword
    containing ".pptx" would never match, which is intended).
    """
    if not keyword:
        return True
    base = os.path.splitext(os.path.basename(pptx_path))[0]
    return keyword in base


def collect_pptx(root, depth, filename_kw=None):
    """Return a sorted list of .pptx paths under root, limited to `depth` levels.

    depth=None means unlimited. depth=0 would mean root only (we treat
    depth via descent count, so callers pass depth=3 for 3 levels down).
    If filename_kw is given, only files whose basename contains it are returned.
    """
    results = []
    root = root.rstrip(os.sep)
    start_level = root.count(os.sep)
    for dirpath, dirnames, filenames in os.walk(root):
        level = dirpath.count(os.sep) - start_level
        if depth is not None and level > depth:
            dirnames[:] = []  # prune: do not descend further
            continue
        for fn in filenames:
            if fn.lower().endswith(".pptx") and filename_matches(fn, filename_kw):
                results.append(os.path.join(dirpath, fn))
    return sorted(results)


def process_directory(root, depth, filename_kw=None):
    files = collect_pptx(root, depth, filename_kw)
    if not files:
        print("No .pptx files found under: %s" % root)
        return
    print("Found %d .pptx file(s) (depth=%s%s) under %s" %
          (len(files), "unlimited" if depth is None else depth,
           (", filename~=%r" % filename_kw) if filename_kw else "", root))
    for p in files:
        process_file(p, None)


def main():
    args = sys.argv[1:]

    # Manual arg parsing to keep -auto / --depth / -filename / optional output flexible.
    auto = False
    depth = 3
    filename_kw = None
    positional = []

    i = 0
    while i < len(args):
        a = args[i]
        if a == "-auto":
            auto = True
        elif a == "--depth":
            i += 1
            if i >= len(args):
                print("Error: --depth requires a number")
                sys.exit(1)
            try:
                depth = int(args[i])
            except ValueError:
                print("Error: --depth must be an integer (0 = unlimited)")
                sys.exit(1)
        elif a.startswith("--depth="):
            try:
                depth = int(a.split("=", 1)[1])
            except ValueError:
                print("Error: --depth must be an integer (0 = unlimited)")
                sys.exit(1)
        elif a == "-filename":
            i += 1
            if i >= len(args):
                print("Error: -filename requires a keyword")
                sys.exit(1)
            filename_kw = args[i]
        elif a.startswith("-filename="):
            filename_kw = a.split("=", 1)[1]
            if filename_kw == "":
                print("Error: -filename requires a non-empty keyword")
                sys.exit(1)
        else:
            positional.append(a)
        i += 1

    if not positional:
        print("Usage:")
        print("  Single file : python3 extract_pptx_notes.py <input.pptx> [output.txt]")
        print("  Directory   : python3 extract_pptx_notes.py -auto <dir> [--depth N] [-filename=<kw>]")
        print("                python3 extract_pptx_notes.py <dir>  (auto-detected)")
        print("  Filter      : -filename=<keyword>  only process .pptx whose name contains <keyword>")
        sys.exit(1)

    path = positional[0]
    output = positional[1] if len(positional) > 1 else None

    is_dir = os.path.isdir(path)

    if is_dir or auto:
        if not is_dir:
            print("Error: -auto requires a directory path, got: %s" % path)
            sys.exit(1)
        d = None if depth == 0 else depth
        process_directory(path, d, filename_kw)
    else:
        if not os.path.isfile(path):
            print("Error: file not found: %s" % path)
            sys.exit(1)
        if not filename_matches(path, filename_kw):
            print("Skipped (filename does not contain %r): %s" %
                  (filename_kw, path))
            sys.exit(0)
        process_file(path, output)


if __name__ == "__main__":
    main()
